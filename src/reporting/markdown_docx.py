"""Turn an AI markdown reply into a Word document.

The Discussion tab uses this for the business description and the
maintainability corrections. Headings, bullets, simple tables and bold
markers are preserved so the file opens as a structured document.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

#: Dewey already writes every other Word deliverable in this folder.
WORD_SUBDIR = "word"

_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
_BULLET_RE = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_RULE_RE = re.compile(r"^\s*-{3,}\s*$")
_INLINE_RE = re.compile(r"(\*\*[^*]+?\*\*|\*[^*\n]+?\*|`[^`]+`)")
_REMAINING_FR = re.compile(r"^#+\s*Reste\s+[àa]\s+faire\s*$", re.IGNORECASE | re.MULTILINE)
_REMAINING_EN = re.compile(r"^#+\s*Remaining work\s*$", re.IGNORECASE | re.MULTILINE)


def sanitize_alias(alias: str) -> str:
    """Make an org alias safe to embed in a Windows file name."""

    cleaned = _INVALID_FILENAME.sub("_", (alias or "").strip())
    cleaned = cleaned.strip(" .")
    return (cleaned or "org")[:80]


def document_filename(kind: str, alias: str, stamp: str) -> str:
    """Return the deliverable name requested for one of the two documents.

    ``stamp`` is the day the run started, ``YYYYMMDD``.
    """

    safe = sanitize_alias(alias)
    if kind == "business":
        return f"Description_metier_org_{safe}_{stamp}.docx"
    return f"Correction_Maintenabilité_metier_org_{safe}_{stamp}.docx"


def word_output_path(output_dir: str | Path, filename: str) -> Path:
    return Path(output_dir) / WORD_SUBDIR / filename


def strip_code_fence(text: str) -> str:
    """Drop a single markdown fence wrapped around a whole reply."""

    stripped = (text or "").strip()
    if not stripped.startswith("```"):
        return stripped
    lines = stripped.splitlines()
    if lines and lines[0].startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].startswith("```"):
        lines = lines[:-1]
    return "\n".join(lines).strip()


def ensure_remaining_section(markdown: str, gaps: list[str], language: str) -> str:
    """Guarantee the business document ends with a remaining-work section."""

    english = (language or "").lower().startswith("en")
    if english:
        if _REMAINING_EN.search(markdown or ""):
            return markdown
        heading = "Remaining work"
        intro = (
            "The following points still have to be confirmed or completed "
            "with the business teams:"
        )
        empty = "- No automatic gap was detected; confirm the scope with the business teams."
    else:
        if _REMAINING_FR.search(markdown or ""):
            return markdown
        heading = "Reste à faire"
        intro = (
            "Les points suivants restent à confirmer ou à compléter avec les équipes métier :"
        )
        empty = "- Aucune lacune automatique n'a été détectée ; à confirmer avec le métier."
    bullets = [f"- {gap}" for gap in gaps if gap.strip()] or [empty]
    block = "\n".join([f"# {heading}", "", intro, *bullets, ""])
    return (markdown or "").rstrip() + "\n\n" + block


def prepare_maintainability_markdown(skeleton: str, reply: str) -> str:
    """Keep the object / family outline even when the model rewrites it.

    The reply is used as-is when every skeleton heading is still present.
    Otherwise the detected outline stays in front and the model text is
    appended, so the document remains grouped by object.
    """

    cleaned = strip_code_fence(reply)
    if not cleaned:
        return skeleton
    required = [line.strip() for line in skeleton.splitlines() if line.startswith("#")]
    present = {line.strip() for line in cleaned.splitlines() if line.startswith("#")}
    if required and all(item in present for item in required):
        return cleaned
    title = "Analyse complémentaire de l'IA"
    if "# Maintainability corrections" in skeleton:
        title = "Additional AI analysis"
    return skeleton.rstrip() + f"\n\n# {title}\n\n" + cleaned + "\n"


def write_markdown_document(
    markdown: str,
    output_path: str | Path,
    *,
    title: str,
    subtitle: str,
    details: str = "",
    language: str = "fr",
) -> Path:
    """Write ``markdown`` to ``output_path`` and return that path."""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure(document)
    document.core_properties.title = title
    document.core_properties.subject = subtitle
    _add_cover(document, title=title, subtitle=subtitle, details=details)
    _add_table_of_contents(document, language)
    _render_markdown(document, strip_code_fence(markdown))
    document.save(output_path)
    return output_path


def _configure(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)


def _add_cover(document: Document, *, title: str, subtitle: str, details: str) -> None:
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = 1
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(26)

    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.alignment = 1
    subtitle_run = subtitle_paragraph.add_run(subtitle)
    subtitle_run.font.size = Pt(14)

    if details.strip():
        details_paragraph = document.add_paragraph()
        details_paragraph.alignment = 1
        for index, line in enumerate(details.strip().splitlines()):
            if index:
                details_paragraph.add_run().add_break()
            run = details_paragraph.add_run(line)
            run.italic = True
            run.font.size = Pt(10)

    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def _add_table_of_contents(document: Document, language: str) -> None:
    english = (language or "").lower().startswith("en")
    document.add_heading("Contents" if english else "Sommaire", level=1)
    hint = document.add_paragraph(
        "In Word, right-click this table and choose Update Field."
        if english
        else "Dans Word, clic droit sur ce sommaire puis Mettre à jour les champs."
    )
    for run in hint.runs:
        run.italic = True

    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _render_markdown(document: Document, markdown: str) -> None:
    lines = (markdown or "").splitlines()
    if not any(line.strip() for line in lines):
        document.add_paragraph("Document vide." if document.core_properties.language != "en" else "Empty document.")
        return

    index = 0
    in_code = False
    code_buffer: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```"):
            if in_code:
                _add_code_block(document, "\n".join(code_buffer))
                code_buffer = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_buffer.append(line)
            index += 1
            continue
        if not line.strip() or _RULE_RE.match(line):
            index += 1
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group(1)), 3)
            document.add_heading(heading.group(2).strip(), level=level)
            index += 1
            continue
        if _is_table_line(line):
            block: list[str] = []
            while index < len(lines) and _is_table_line(lines[index]):
                block.append(lines[index])
                index += 1
            _add_table(document, block)
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            _add_list_item(document, bullet.group(1), numbered=False)
            index += 1
            continue
        number = _NUMBER_RE.match(line)
        if number:
            _add_list_item(document, number.group(1), numbered=True)
            index += 1
            continue

        paragraph = document.add_paragraph()
        _add_inline(paragraph, line.strip())
        index += 1

    if code_buffer:
        _add_code_block(document, "\n".join(code_buffer))


def _add_inline(paragraph, text: str) -> None:
    position = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _add_list_item(document: Document, text: str, *, numbered: bool) -> None:
    style = "List Number" if numbered else "List Bullet"
    try:
        paragraph = document.add_paragraph(style=style)
    except (KeyError, ValueError):
        paragraph = document.add_paragraph()
        text = ("1. " if numbered else "• ") + text
    _add_inline(paragraph, text.strip())


def _add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    parts = text.split("\n")
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if index < len(parts) - 1:
            run.add_break()


def _is_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_separator(cells: list[str]) -> bool:
    if not cells:
        return False
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def _add_table(document: Document, rows: list[str]) -> None:
    parsed = [_split_row(row) for row in rows]
    parsed = [row for row in parsed if not _is_separator(row)]
    if not parsed:
        return
    width = max(len(row) for row in parsed)
    table = document.add_table(rows=len(parsed), cols=width)
    try:
        table.style = "Table Grid"
    except (KeyError, ValueError):
        pass
    for row_index, row in enumerate(parsed):
        for column in range(width):
            value = row[column] if column < len(row) else ""
            cell = table.cell(row_index, column)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _add_inline(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
