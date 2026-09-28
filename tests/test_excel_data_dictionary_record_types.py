"""The Data Dictionary workbook lists record types on a sheet after "Synthese"."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.core.models import (
    FieldInfo,
    ObjectInfo,
    RecordTypeInfo,
    RecordTypeVisibility,
    SecurityArtifact,
)
from src.parsers.salesforce_parser import SalesforceMetadataParser
from src.reporting.excel_writer import ExcelReportWriter

RECORD_TYPE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<RecordType xmlns="http://soap.sforce.com/2006/04/metadata">
    <fullName>NewBusiness</fullName>
    <active>true</active>
    <businessProcess>SalesProcess</businessProcess>
    <compactLayoutAssignment>OppCompact</compactLayoutAssignment>
    <description>Nouvelles affaires</description>
    <label>New Business</label>
    <picklistValues>
        <picklist>StageName</picklist>
        <values><fullName>Closed%20Won</fullName><default>false</default></values>
        <values><fullName>Prospecting</fullName><default>true</default></values>
    </picklistValues>
</RecordType>
"""


def _rows(sheet) -> list[dict]:
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, (cell.value for cell in row))) for row in sheet.iter_rows(min_row=2)]


def test_parser_reads_business_process_layout_and_picklists(tmp_path: Path):
    folder = tmp_path / "force-app" / "main" / "default" / "objects" / "Opportunity" / "recordTypes"
    folder.mkdir(parents=True)
    (folder / "NewBusiness.recordType-meta.xml").write_text(RECORD_TYPE_XML, encoding="utf-8")

    snapshot = SalesforceMetadataParser(tmp_path).parse()
    opportunity = next(obj for obj in snapshot.objects if obj.api_name == "Opportunity")
    record_type = opportunity.record_types[0]

    assert record_type.business_process == "SalesProcess"
    assert record_type.compact_layout == "OppCompact"
    assert record_type.picklist_values == {"StageName": ["Closed Won", "Prospecting"]}


def test_record_types_sheet_follows_synthese_with_visibility(tmp_path: Path):
    opportunity = ObjectInfo(
        api_name="Opportunity",
        label="Opportunite",
        fields=[FieldInfo(api_name="StageName")],
        record_types=[
            RecordTypeInfo(
                full_name="NewBusiness",
                label="New Business",
                description="Nouvelles affaires",
                active=True,
                business_process="SalesProcess",
                picklist_values={"StageName": ["Prospecting", "Closed Won"]},
            )
        ],
    )
    profile = SecurityArtifact(
        name="Sales",
        kind="profile",
        record_type_visibilities=[
            RecordTypeVisibility(record_type="Opportunity.NewBusiness", visible=True, default=True)
        ],
    )
    permission_set = SecurityArtifact(
        name="Renewals",
        kind="permission_set",
        record_type_visibilities=[
            RecordTypeVisibility(record_type="Opportunity.NewBusiness", visible=True)
        ],
    )

    paths = ExcelReportWriter().write_data_dictionary_workbooks(
        [opportunity], tmp_path, profiles=[profile], permission_sets=[permission_set]
    )
    workbook = load_workbook(paths[0])

    assert workbook.sheetnames == ["Synthese", "Record Types", "Opportunity"]
    [summary_row] = _rows(workbook["Synthese"])
    assert summary_row["Record Type"] == "1 : New Business"
    assert workbook["Synthese"]["J1"].value == "Nb record types"
    assert summary_row["Nb record types"] == "Actif : 1 | Inactif : 0"
    [row] = _rows(workbook["Record Types"])
    assert row["Objet (API Name)"] == "Opportunity"
    assert row["Record Type (API Name)"] == "NewBusiness"
    assert row["Label"] == "New Business"
    assert row["Actif"] == "Oui"
    assert row["Business Process"] == "SalesProcess"
    assert row["Valeurs de picklist disponibles"] == "StageName : Prospecting, Closed Won"
    assert row["Profils (visible)"] == "Sales"
    assert row["Profils (par defaut)"] == "Sales"
    assert row["Permission Sets (visible)"] == "Renewals"


def test_synthese_record_type_column_lists_sorted_labels():
    obj = ObjectInfo(
        api_name="Case",
        record_types=[
            RecordTypeInfo(full_name="Support", label="Support"),
            RecordTypeInfo(full_name="Claim", label="Reclamation"),
            RecordTypeInfo(full_name="NoLabel"),
        ],
    )
    assert ExcelReportWriter._record_types_summary(obj) == "3 : NoLabel | Reclamation | Support"
    assert ExcelReportWriter._record_types_summary(ObjectInfo(api_name="Account")) == ""
    assert ExcelReportWriter._record_types_activity(obj) == "Actif : 0 | Inactif : 3"


def test_record_types_sheet_without_record_types_or_security(tmp_path: Path):
    account = ObjectInfo(api_name="Account", fields=[FieldInfo(api_name="Name")])

    paths = ExcelReportWriter().write_data_dictionary_workbooks([account], tmp_path)
    sheet = load_workbook(paths[0])["Record Types"]

    headers = [cell.value for cell in sheet[1]]
    assert "Profils (visible)" not in headers
    assert sheet.cell(row=2, column=1).value.startswith("Aucun record type")
