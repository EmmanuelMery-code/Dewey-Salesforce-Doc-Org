"""Tests for the "En conception" objects and fields of the Data Dictionary."""

from __future__ import annotations

import pytest

from src.core.data_dictionary_selection import (
    DELIVERED_STATUS,
    IN_DESIGN_STATUS,
    DataDictionarySelection,
)
from src.core.models import FieldInfo, ObjectInfo, RelationshipInfo
from src.ui.data_dictionary_screen import virtual_entries
from src.ui.data_dictionary_screen.field_info import _DataDictionaryFieldInfoMixin
from src.ui.data_dictionary_screen.virtual_entries import (
    _DataDictionaryVirtualEntriesMixin,
)

STATUS_OPTIONS = ["-", "en dév.", DELIVERED_STATUS, IN_DESIGN_STATUS]


class _App:
    def __init__(self) -> None:
        self.settings: dict[str, object] = {}
        self.save_count = 0

    def _save_settings(self) -> None:
        self.save_count += 1

    def _t(self, key: str, **kwargs: object) -> str:
        if not kwargs:
            return key
        rendered = ", ".join(f"{name}={value}" for name, value in sorted(kwargs.items()))
        return f"{key}[{rendered}]"


class _Screen(_DataDictionaryFieldInfoMixin, _DataDictionaryVirtualEntriesMixin):
    """Stand-in for the Data Dictionary screen, without a Tk display.

    ``metadata_fields`` plays the role of the field metadata files that
    ``_real_object_fields`` normally reads from disk.
    """

    STATUS_OPTIONS = STATUS_OPTIONS

    def __init__(
        self,
        metadata_objects: set[str] | None = None,
        metadata_fields: dict[str, list[str]] | None = None,
    ) -> None:
        self.app = _App()
        self._metadata_objects = metadata_objects or set()
        self._metadata_fields = metadata_fields or {}
        self.all_objects = sorted(self._metadata_objects)
        self.selected_objects: set[str] = set()
        self.current_comment_object: str | None = None
        self.current_comment_field: str | None = None
        self.object_comments: dict[str, str] = {}
        self.object_piloted_by: dict[str, str] = {}
        self.object_status: dict[str, str] = {}
        self.object_squad: dict[str, str] = {}
        self.object_squad_consumer: dict[str, str] = {}
        self.field_comments: dict[str, dict[str, str]] = {}
        self.field_piloted_by: dict[str, dict[str, str]] = {}
        self.field_status: dict[str, dict[str, str]] = {}
        self.virtual_objects: dict[str, dict[str, str]] = {}
        self.virtual_fields: dict[str, dict[str, dict[str, str]]] = {}

    def _real_object_fields(self, obj: str) -> list[tuple[str, str]]:
        return [(name, name) for name in self._metadata_fields.get(obj, [])]

    def _persist_comments(self) -> None:
        self.app.settings["dd_object_status"] = dict(self.object_status)
        self.app._save_settings()

    def _persist_field_comments(self) -> None:
        self.app.settings["dd_field_status"] = {
            obj: dict(fields) for obj, fields in self.field_status.items()
        }
        self.app._save_settings()


class TestCreation:
    def test_a_virtual_object_is_selected_and_forced_in_design(self) -> None:
        screen = _Screen(metadata_objects={"Account"})

        screen._create_virtual_object("Projet__c", "Projet")

        assert screen.virtual_objects == {"Projet__c": {"label": "Projet"}}
        assert screen.object_status["Projet__c"] == IN_DESIGN_STATUS
        assert "Projet__c" in screen.selected_objects
        assert screen.all_objects == ["Account", "Projet__c"]

    def test_a_virtual_field_is_forced_in_design(self) -> None:
        screen = _Screen(metadata_objects={"Account"})

        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency")

        assert screen.virtual_fields == {
            "Account": {"Budget__c": {"label": "Budget", "type": "Currency"}}
        }
        assert screen.field_status["Account"]["Budget__c"] == IN_DESIGN_STATUS

    def test_a_virtual_field_shows_up_in_the_object_field_list(self) -> None:
        screen = _Screen(
            metadata_objects={"Account"}, metadata_fields={"Account": ["Name"]}
        )

        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency")

        assert screen._list_object_fields("Account") == [
            ("Budget__c", "Budget"),
            ("Name", "Name"),
        ]

    def test_a_comment_typed_at_creation_is_stored(self) -> None:
        screen = _Screen(metadata_objects={"Account"})

        screen._create_virtual_object("Projet__c", "Projet", "Objet du chantier 2027")
        screen._create_virtual_field(
            "Account", "Budget__c", "Budget", "Currency", "Montant previsionnel"
        )

        assert screen.object_comments == {"Projet__c": "Objet du chantier 2027"}
        assert screen.field_comments == {"Account": {"Budget__c": "Montant previsionnel"}}

    def test_no_comment_stores_nothing_at_all(self) -> None:
        """Blank leaves the mappings untouched, as the save panels do."""
        screen = _Screen(metadata_objects={"Account"})

        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency", "")

        assert screen.object_comments == {}
        assert screen.field_comments == {}


class TestPersistence:
    def test_both_mappings_round_trip_through_the_settings(self) -> None:
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Projet__c", "Budget__c", "Budget", "Currency")

        selection = DataDictionarySelection.from_settings(screen.app.settings)

        assert screen.app.settings["dd_virtual_objects"] == {
            "Projet__c": {"label": "Projet"}
        }
        assert screen.app.settings["dd_virtual_fields"] == {
            "Projet__c": {"Budget__c": {"label": "Budget", "type": "Currency"}}
        }
        assert selection.virtual_objects == {"Projet__c": {"label": "Projet"}}
        assert selection.virtual_fields == {
            "Projet__c": {"Budget__c": {"label": "Budget", "type": "Currency"}}
        }

    def test_corrupt_settings_are_ignored(self) -> None:
        selection = DataDictionarySelection.from_settings(
            {"dd_virtual_objects": "nope", "dd_virtual_fields": []}
        )

        assert selection.virtual_objects == {}
        assert selection.virtual_fields == {}

    def test_deleting_a_virtual_object_drops_its_fields_and_statuses(self) -> None:
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Projet__c", "Budget__c", "Budget", "Currency")

        screen._delete_virtual_object("Projet__c")

        assert screen.virtual_objects == {}
        assert screen.virtual_fields == {}
        assert screen.object_status == {}
        assert screen.field_status == {}
        assert screen.selected_objects == set()
        assert screen.all_objects == ["Account"]

    def test_deleting_a_virtual_field_leaves_the_object_alone(self) -> None:
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency")

        screen._delete_virtual_field("Account", "Budget__c")

        assert screen.virtual_fields == {"Account": {}}
        assert screen.field_status == {"Account": {}}


class TestValidation:
    def test_an_empty_api_name_or_label_is_rejected(self) -> None:
        screen = _Screen()

        assert (
            screen._virtual_object_error("", "Projet")
            == "data_dictionary_virtual_error_api_name_required"
        )
        assert (
            screen._virtual_object_error("Projet__c", "")
            == "data_dictionary_virtual_error_label_required"
        )

    def test_whitespace_inside_the_api_name_is_rejected(self) -> None:
        screen = _Screen()

        assert (
            screen._virtual_object_error("Projet c", "Projet")
            == "data_dictionary_virtual_error_api_name_whitespace"
        )

    def test_a_duplicated_virtual_object_is_rejected(self) -> None:
        screen = _Screen()
        screen._create_virtual_object("Projet__c", "Projet")

        assert screen._virtual_object_error("Projet__c", "Projet") == (
            "data_dictionary_virtual_error_object_duplicate[name=Projet__c]"
        )

    def test_an_object_already_in_the_metadata_is_rejected(self) -> None:
        screen = _Screen(metadata_objects={"Account"})

        assert screen._virtual_object_error("Account", "Compte") == (
            "data_dictionary_virtual_error_object_exists[name=Account]"
        )

    def test_a_missing_field_type_is_rejected(self) -> None:
        screen = _Screen(metadata_objects={"Account"})

        assert (
            screen._virtual_field_error("Account", "Budget__c", "Budget", "")
            == "data_dictionary_virtual_error_type_required"
        )

    def test_a_duplicated_virtual_field_is_rejected(self) -> None:
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency")

        assert screen._virtual_field_error(
            "Account", "Budget__c", "Budget", "Currency"
        ) == ("data_dictionary_virtual_error_field_duplicate[name=Budget__c]")

    def test_a_field_already_in_the_metadata_is_rejected(self) -> None:
        screen = _Screen(
            metadata_objects={"Account"}, metadata_fields={"Account": ["Name"]}
        )

        assert screen._virtual_field_error("Account", "Name", "Nom", "Text") == (
            "data_dictionary_virtual_error_field_exists[name=Name]"
        )

    def test_a_valid_declaration_is_accepted(self) -> None:
        screen = _Screen(
            metadata_objects={"Account"}, metadata_fields={"Account": ["Name"]}
        )

        assert screen._virtual_object_error("Projet__c", "Projet") is None
        assert screen._virtual_field_error("Account", "Budget__c", "Budget", "Currency") is None


class TestPromotion:
    def test_a_virtual_object_found_in_the_metadata_is_delivered(self) -> None:
        screen = _Screen()
        screen._create_virtual_object("Projet__c", "Projet")
        screen._metadata_objects = {"Projet__c"}

        promoted_objects, promoted_fields = screen._promote_virtual_entries()

        assert promoted_objects == ["Projet__c"]
        assert promoted_fields == {}
        assert screen.virtual_objects == {}
        assert screen.object_status["Projet__c"] == DELIVERED_STATUS
        assert screen.app.settings["dd_object_status"] == {"Projet__c": DELIVERED_STATUS}

    def test_a_virtual_field_found_in_the_metadata_is_delivered(self) -> None:
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency")
        screen._metadata_fields = {"Account": ["Name", "Budget__c"]}

        promoted_objects, promoted_fields = screen._promote_virtual_entries()

        assert promoted_objects == []
        assert promoted_fields == {"Account": ["Budget__c"]}
        assert screen.virtual_fields == {"Account": {}}
        assert screen.field_status["Account"]["Budget__c"] == DELIVERED_STATUS
        assert screen.app.settings["dd_field_status"] == {
            "Account": {"Budget__c": DELIVERED_STATUS}
        }

    def test_the_fields_of_a_promoted_object_are_promoted_in_the_same_pass(self) -> None:
        screen = _Screen()
        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Projet__c", "Budget__c", "Budget", "Currency")
        screen._metadata_objects = {"Projet__c"}
        screen._metadata_fields = {"Projet__c": ["Budget__c"]}

        promoted_objects, promoted_fields = screen._promote_virtual_entries()

        assert promoted_objects == ["Projet__c"]
        assert promoted_fields == {"Projet__c": ["Budget__c"]}
        assert screen.object_status["Projet__c"] == DELIVERED_STATUS
        assert screen.field_status["Projet__c"]["Budget__c"] == DELIVERED_STATUS

    def test_promotion_overrides_a_manually_set_status(self) -> None:
        """The metadata is ground truth, whatever the user picked meanwhile."""
        screen = _Screen()
        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Projet__c", "Budget__c", "Budget", "Currency")
        screen.object_status["Projet__c"] = "en dév."
        screen.field_status["Projet__c"]["Budget__c"] = "en dév."
        screen._metadata_objects = {"Projet__c"}
        screen._metadata_fields = {"Projet__c": ["Budget__c"]}

        screen._promote_virtual_entries()

        assert screen.object_status["Projet__c"] == DELIVERED_STATUS
        assert screen.field_status["Projet__c"]["Budget__c"] == DELIVERED_STATUS

    def test_entries_still_absent_from_the_metadata_stay_in_design(self) -> None:
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Account", "Budget__c", "Budget", "Currency")

        assert screen._promote_virtual_entries() == ([], {})
        assert screen.virtual_objects == {"Projet__c": {"label": "Projet"}}
        assert screen.object_status["Projet__c"] == IN_DESIGN_STATUS
        assert screen.field_status["Account"]["Budget__c"] == IN_DESIGN_STATUS


class TestPromotionRecap:
    def test_nothing_promoted_shows_no_dialog(self, monkeypatch: pytest.MonkeyPatch) -> None:
        shown: list[tuple[str, str]] = []
        monkeypatch.setattr(
            virtual_entries.messagebox,
            "showinfo",
            lambda title, message: shown.append((title, message)),
        )
        screen = _Screen(metadata_objects={"Account"})
        screen._create_virtual_object("Projet__c", "Projet")

        screen._promote_virtual_entries_and_report()

        assert shown == []

    def test_promoted_entries_are_recapped_grouped_by_object(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        shown: list[tuple[str, str]] = []
        monkeypatch.setattr(
            virtual_entries.messagebox,
            "showinfo",
            lambda title, message: shown.append((title, message)),
        )
        screen = _Screen()
        screen._create_virtual_object("Projet__c", "Projet")
        screen._create_virtual_field("Projet__c", "Budget__c", "Budget", "Currency")
        screen._create_virtual_field("Account", "Segment__c", "Segment", "Picklist")
        screen._metadata_objects = {"Account", "Projet__c"}
        screen._metadata_fields = {
            "Account": ["Segment__c"],
            "Projet__c": ["Budget__c"],
        }

        screen._promote_virtual_entries_and_report()

        assert len(shown) == 1
        title, message = shown[0]
        assert title == "data_dictionary_virtual_promoted_title"
        assert "- Account\n    * Segment__c" in message
        assert (
            "- Projet__cdata_dictionary_virtual_promoted_object_suffix\n    * Budget__c"
            in message
        )


class TestSelectionSynthesis:
    def test_a_virtual_object_and_its_fields_are_synthesized(self) -> None:
        selection = DataDictionarySelection(
            objects={"Projet__c"},
            virtual_objects={"Projet__c": {"label": "Projet"}},
            virtual_fields={
                "Projet__c": {"Budget__c": {"label": "Budget", "type": "Currency"}}
            },
        )

        objects = selection.apply([ObjectInfo(api_name="Account", label="Account")])

        assert [obj.api_name for obj in objects] == ["Projet__c"]
        projet = objects[0]
        assert projet.label == "Projet"
        assert projet.custom is True
        assert projet.dewey_status == IN_DESIGN_STATUS
        budget = projet.fields[0]
        assert (budget.api_name, budget.label, budget.data_type) == (
            "Budget__c",
            "Budget",
            "Currency",
        )
        assert budget.dewey_status == IN_DESIGN_STATUS

    def test_a_virtual_field_is_appended_to_a_real_object(self) -> None:
        selection = DataDictionarySelection(
            objects={"Account"},
            virtual_fields={
                "Account": {"Segment__c": {"label": "Segment", "type": "Picklist"}}
            },
        )

        account = selection.apply(
            [
                ObjectInfo(
                    api_name="Account",
                    label="Account",
                    fields=[FieldInfo(api_name="Name", label="Name")],
                )
            ]
        )[0]

        assert [field.api_name for field in account.fields] == ["Name", "Segment__c"]
        assert account.fields[1].dewey_status == IN_DESIGN_STATUS

    def test_an_unselected_virtual_object_is_not_synthesized(self) -> None:
        selection = DataDictionarySelection(
            objects=set(), virtual_objects={"Projet__c": {"label": "Projet"}}
        )

        assert selection.apply([]) == []

    def test_a_virtual_entry_already_parsed_is_not_duplicated(self) -> None:
        """Belt and braces: promotion normally removes it first."""
        selection = DataDictionarySelection(
            objects={"Account"},
            virtual_objects={"Account": {"label": "Compte"}},
            virtual_fields={"Account": {"Name": {"label": "Nom", "type": "Text"}}},
        )

        objects = selection.apply(
            [
                ObjectInfo(
                    api_name="Account",
                    label="Account",
                    fields=[FieldInfo(api_name="Name", label="Name")],
                )
            ]
        )

        assert len(objects) == 1
        assert [field.api_name for field in objects[0].fields] == ["Name"]


class TestDataModelGraphExclusion:
    """No edge may survive the shape it was attached to."""

    @staticmethod
    def _links(objects: list[ObjectInfo]) -> list[tuple[str, str, str]]:
        from src.core.data_model_graph import plan_data_model_tabs

        return [
            (link.source, link.target, link.field_name)
            for cluster in plan_data_model_tabs(objects)
            for link in cluster.links
        ]

    def test_a_relationship_targeting_a_virtual_object_is_dropped(self) -> None:
        objects = [
            ObjectInfo(
                api_name="Account",
                label="Account",
                fields=[FieldInfo(api_name="Name", label="Name")],
                relationships=[
                    RelationshipInfo(
                        field_name="ProjetId__c",
                        relationship_type="Lookup",
                        targets=["Projet__c"],
                    )
                ],
            ),
            ObjectInfo(api_name="Projet__c", label="Projet", is_virtual=True),
        ]

        from src.core.data_model_graph import plan_data_model_tabs

        clusters = plan_data_model_tabs(objects)

        assert self._links(objects) == []
        assert [name for cluster in clusters for name in cluster.object_names] == [
            "Account"
        ]

    def test_a_relationship_carried_by_a_virtual_field_is_dropped(self) -> None:
        objects = [
            ObjectInfo(
                api_name="Account",
                label="Account",
                fields=[
                    FieldInfo(api_name="Name", label="Name"),
                    FieldInfo(api_name="AtelierId__c", label="Atelier", is_virtual=True),
                ],
                relationships=[
                    RelationshipInfo(
                        field_name="AtelierId__c",
                        relationship_type="Lookup",
                        targets=["Contact"],
                    )
                ],
            ),
            ObjectInfo(api_name="Contact", label="Contact"),
        ]

        assert self._links(objects) == []

    def test_a_real_relationship_between_real_objects_survives(self) -> None:
        objects = [
            ObjectInfo(
                api_name="Account",
                label="Account",
                dewey_status=IN_DESIGN_STATUS,
                fields=[FieldInfo(api_name="ParentId", label="Parent")],
                relationships=[
                    RelationshipInfo(
                        field_name="ParentId",
                        relationship_type="Lookup",
                        targets=["Contact"],
                    )
                ],
            ),
            ObjectInfo(api_name="Contact", label="Contact"),
        ]

        assert self._links(objects) == [("Account", "Contact", "ParentId")]


class TestGeneratedDeliverables:
    """Virtual entries belong to the Excel and Word dictionaries only.

    The fixture covers the cases that matter: a virtual object with a virtual
    field, a virtual object with no field at all, a virtual field on a real
    object, a real object the user manually set to "En conception" — which
    must keep behaving like any other real object — and a real relationship
    between two real objects, so the draw.io edges have something to draw.
    """

    @staticmethod
    def _selected_objects() -> list[ObjectInfo]:
        selection = DataDictionarySelection(
            objects={"Account", "Compte_Parent__c", "Projet__c", "Atelier__c"},
            object_status={"Account": IN_DESIGN_STATUS},
            object_comments={"Projet__c": "Chantier de refonte 2027"},
            field_comments={"Projet__c": {"Budget__c": "Montant previsionnel annuel"}},
            virtual_objects={
                "Projet__c": {"label": "Projet en conception"},
                "Atelier__c": {"label": "Atelier sans champ"},
            },
            virtual_fields={
                "Account": {"Segment__c": {"label": "Segment", "type": "Picklist"}},
                "Projet__c": {"Budget__c": {"label": "Budget", "type": "Currency"}},
            },
        )
        return selection.apply(
            [
                ObjectInfo(
                    api_name="Account",
                    label="Account",
                    fields=[
                        FieldInfo(api_name="Name", label="Name", data_type="Text"),
                        FieldInfo(api_name="ParentId", label="Parent", data_type="Lookup"),
                    ],
                    relationships=[
                        RelationshipInfo(
                            field_name="ParentId",
                            relationship_type="Lookup",
                            targets=["Compte_Parent__c"],
                        )
                    ],
                ),
                ObjectInfo(
                    api_name="Compte_Parent__c",
                    label="Compte parent",
                    fields=[FieldInfo(api_name="Nom__c", label="Nom", data_type="Text")],
                ),
            ]
        )

    @staticmethod
    def _snapshot(tmp_path):
        from src.core.models import MetadataSnapshot

        return MetadataSnapshot(
            source_dir=tmp_path,
            package_roots=[],
            objects=TestGeneratedDeliverables._selected_objects(),
        )

    def test_the_excel_dictionary_carries_them_with_their_status(self, tmp_path) -> None:
        openpyxl = pytest.importorskip("openpyxl")
        from src.reporting.excel_writer import ExcelReportWriter

        paths = ExcelReportWriter().write_data_dictionary_workbooks(
            self._selected_objects(), tmp_path, filename_base="dd"
        )

        workbook = openpyxl.load_workbook(paths[0])
        summary = list(workbook["Synthese"].values)
        headers = list(summary[0])
        rows = {row[0]: row for row in summary[1:]}
        assert rows["Projet__c"][1] == "Projet en conception"
        assert rows["Projet__c"][headers.index("Status")] == IN_DESIGN_STATUS
        assert (
            rows["Projet__c"][headers.index("Commentaire Dewey")]
            == "Chantier de refonte 2027"
        )

        field_headers = [cell.value for cell in workbook["Account"][1]]
        account_rows = {row[0].value: row for row in workbook["Account"].iter_rows(min_row=2)}
        segment = account_rows["Segment__c"]
        assert segment[field_headers.index("Type")].value == "Picklist"
        assert segment[field_headers.index("Status")].value == IN_DESIGN_STATUS
        projet_rows = {row[0].value: row for row in workbook["Projet__c"].iter_rows(min_row=2)}
        budget = projet_rows["Budget__c"]
        assert budget[field_headers.index("Status")].value == IN_DESIGN_STATUS
        assert (
            budget[field_headers.index("Commentaire Dewey")].value
            == "Montant previsionnel annuel"
        )

    def test_a_field_less_virtual_object_gets_an_empty_excel_sheet(self, tmp_path) -> None:
        openpyxl = pytest.importorskip("openpyxl")
        from src.reporting.excel_writer import ExcelReportWriter

        paths = ExcelReportWriter().write_data_dictionary_workbooks(
            self._selected_objects(), tmp_path, filename_base="dd"
        )

        workbook = openpyxl.load_workbook(paths[0])
        summary = list(workbook["Synthese"].values)
        headers = list(summary[0])
        rows = {row[0]: row for row in summary[1:]}
        assert rows["Atelier__c"][1] == "Atelier sans champ"
        assert rows["Atelier__c"][headers.index("Status")] == IN_DESIGN_STATUS
        assert rows["Atelier__c"][headers.index("Nb champs")] == 0

        sheet = workbook["Atelier__c"]
        assert [cell.value for cell in sheet[1]] == [
            cell.value for cell in workbook["Account"][1]
        ]
        assert sheet.cell(row=2, column=2).value is None
        assert (
            sheet.cell(row=2, column=1).value
            == "Aucun champ en conception declare pour cet objet."
        )

    def test_a_real_object_with_no_field_is_still_skipped_in_excel(self, tmp_path) -> None:
        """Long-standing behaviour, unrelated to the design entries."""
        openpyxl = pytest.importorskip("openpyxl")
        from src.reporting.excel_writer import ExcelReportWriter

        paths = ExcelReportWriter().write_data_dictionary_workbooks(
            [ObjectInfo(api_name="Empty__c", label="Empty")], tmp_path, filename_base="dd"
        )

        workbook = openpyxl.load_workbook(paths[0])
        assert workbook.sheetnames == ["Synthese"]

    def test_the_html_dictionary_shows_none_of_them(self, tmp_path) -> None:
        from src.reporting.html_writer import HtmlReportWriter

        writer = HtmlReportWriter(tmp_path)
        output_path = tmp_path / "html" / "dd.html"
        writer.write_combined_data_dictionary_html(self._snapshot(tmp_path), output_path)

        content = output_path.read_text(encoding="utf-8")
        for absent in (
            "Projet__c",
            "Atelier__c",
            "Budget__c",
            "Segment__c",
            "Chantier de refonte 2027",
            "Montant previsionnel annuel",
        ):
            assert absent not in content
        # A real object manually set to "En conception" keeps its page.
        assert "Account" in content
        assert "Name" in content
        assert f"<strong>Status:</strong> {IN_DESIGN_STATUS}" in content
        # The "Champs" card counts what the Fields tab actually lists: the two
        # real fields of Account, not the virtual Segment__c.
        assert "<span>Champs</span><span class=\"value\">2</span>" in content

    def test_the_html_object_pages_skip_the_virtual_objects(self, tmp_path) -> None:
        from src.reporting.html_writer import HtmlReportWriter

        pages = HtmlReportWriter(tmp_path).write_object_pages(self._snapshot(tmp_path))

        assert set(pages) == {"Account", "Compte_Parent__c"}
        content = pages["Account"].read_text(encoding="utf-8")
        assert "Segment__c" not in content

    def test_the_word_dictionary_carries_them_with_their_status(self, tmp_path) -> None:
        docx = pytest.importorskip("docx")
        from src.reporting.word_writer import WordReportWriter

        output_path = WordReportWriter().write_data_dictionary_document(
            self._snapshot(tmp_path), tmp_path / "dd.docx"
        )

        document = docx.Document(output_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
        assert "Projet__c" in text
        assert "Projet en conception" in text
        assert IN_DESIGN_STATUS in text
        assert "Segment__c" in text
        assert "Budget__c" in text
        # The object comment reaches its information table. The field comment
        # cannot: the Word fields table renders Label / API Name / Type /
        # Description only, with no per-field Dewey column — a pre-existing
        # limitation shared with real fields, Excel being the only deliverable
        # that carries per-field Dewey values.
        assert "Chantier de refonte 2027" in text
        assert "Montant previsionnel annuel" not in text
        assert [cell.text for cell in self._chapter_tables(document, "Projet__c")[-1].rows[0].cells] == [
            "Label",
            "API Name",
            "Type",
            "Description",
        ]

    @staticmethod
    def _chapter_tables(document, chapter_fragment: str) -> list:
        """The tables of the object chapter whose title contains
        ``chapter_fragment``, in document order."""
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        chapter = ""
        tables = []
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                if paragraph.style.name.startswith("Heading 1"):
                    chapter = paragraph.text
            elif child.tag.endswith("}tbl") and chapter_fragment in chapter:
                tables.append(Table(child, document))
        return tables

    def test_a_field_less_virtual_object_gets_an_empty_word_table(self, tmp_path) -> None:
        docx = pytest.importorskip("docx")
        from src.reporting.word_writer import WordReportWriter

        output_path = WordReportWriter().write_data_dictionary_document(
            self._snapshot(tmp_path), tmp_path / "dd.docx"
        )

        document = docx.Document(output_path)
        assert any(
            "Atelier sans champ" in paragraph.text for paragraph in document.paragraphs
        )

        # The fields table closes the chapter: header row on its own, no data
        # row, and the same 4-column grid as a populated object.
        fields_table = self._chapter_tables(document, "Atelier__c")[-1]
        populated = self._chapter_tables(document, "Projet__c")[-1]
        assert len(fields_table.rows) == 1
        assert len(fields_table.columns) == 4
        assert [cell.text for cell in fields_table.rows[0].cells] == [
            cell.text for cell in populated.rows[0].cells
        ]
        assert len(populated.rows) == 2

    def test_the_drawio_diagram_shows_none_of_them(self, tmp_path) -> None:
        from src.reporting.drawio_writer import DrawioDiagramWriter

        output_path = DrawioDiagramWriter().write_data_model_diagram(
            self._selected_objects(), tmp_path / "data_model.drawio"
        )

        content = output_path.read_text(encoding="utf-8")
        for absent in (
            "Projet__c",
            "Atelier__c",
            "Budget__c",
            "Segment__c",
            "Chantier de refonte 2027",
            "Montant previsionnel annuel",
        ):
            assert absent not in content
        # The real objects and the real field carrying their relationship stay.
        assert "Account" in content
        assert "Compte_Parent__c" in content
        assert "ParentId" in content
        # The header counts are computed after filtering.
        assert "2 objet(s), 1 relation(s)" in content

    def test_a_real_object_with_no_field_is_still_skipped_in_word(self, tmp_path) -> None:
        """Long-standing behaviour, unrelated to the design entries."""
        docx = pytest.importorskip("docx")
        from src.core.models import MetadataSnapshot
        from src.reporting.word_writer import WordReportWriter

        snapshot = MetadataSnapshot(
            source_dir=tmp_path,
            package_roots=[],
            objects=[ObjectInfo(api_name="Empty__c", label="Empty")],
        )
        output_path = WordReportWriter().write_data_dictionary_document(
            snapshot, tmp_path / "dd.docx"
        )

        document = docx.Document(output_path)
        assert not any(
            "Empty__c" in paragraph.text for paragraph in document.paragraphs
        )
