"""Tests for the "dette technique et entorse" workbook.

It is written by every documentation run and can be refreshed on its own
from the technical debt management screen.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import openpyxl

from src.core.models import DeviationItem, MetadataSnapshot, TechnicalDebtItem
from src.core.orchestrator import GenerationResult, SalesforceDocumentationGenerator
from src.reporting.excel_writer import ExcelReportWriter
from src.reporting.excel_writer_debt import (
    DEBT_WORKBOOK_NAME,
    DEVIATIONS_SHEET,
    TECHNICAL_DEBT_SHEET,
)
from src.ui.debt_excel_export import export_debt_workbook

_DEBT_ITEM = TechnicalDebtItem(
    label="Trigger sans handler",
    date_creation="2026-01-05",
    date_resolution="",
    accepted_solution="Laisse en l'etat pour la mise en production",
    target_solution="Extraire la logique dans un handler",
)
_DEVIATION_ITEM = DeviationItem(
    label="Champ en dur dans une formule",
    date_creation="2026-02-10",
    explanation="Contournement valide en comite d'architecture",
)


def _snapshot(tmp_path: Path) -> MetadataSnapshot:
    return MetadataSnapshot(
        source_dir=tmp_path / "source",
        package_roots=[],
        technical_debt=[_DEBT_ITEM],
        deviations=[_DEVIATION_ITEM],
    )


class TestDebtWorkbook:
    def test_the_two_tabs_carry_the_debt_and_the_deviations(self, tmp_path: Path) -> None:
        target = ExcelReportWriter().write_debt_workbook(
            [("MHINT", _DEBT_ITEM)],
            [("MHINT", _DEVIATION_ITEM)],
            tmp_path / DEBT_WORKBOOK_NAME,
        )

        workbook = openpyxl.load_workbook(target)
        assert workbook.sheetnames == [TECHNICAL_DEBT_SHEET, DEVIATIONS_SHEET]

        debt = workbook[TECHNICAL_DEBT_SHEET]
        assert debt["A1"].value == "Libelle"
        assert debt["A2"].value == _DEBT_ITEM.label
        assert debt["E2"].value == _DEBT_ITEM.target_solution

        deviations = workbook[DEVIATIONS_SHEET]
        assert deviations["A2"].value == _DEVIATION_ITEM.label
        assert deviations["C2"].value == _DEVIATION_ITEM.explanation

    def test_the_alias_column_is_added_on_demand(self, tmp_path: Path) -> None:
        """An export covering several orgs must say which one a row is from."""
        target = ExcelReportWriter().write_debt_workbook(
            [("MHINT", _DEBT_ITEM)],
            [("PROD", _DEVIATION_ITEM)],
            tmp_path / DEBT_WORKBOOK_NAME,
            include_alias=True,
        )

        workbook = openpyxl.load_workbook(target)
        assert workbook[TECHNICAL_DEBT_SHEET]["A2"].value == "MHINT"
        assert workbook[DEVIATIONS_SHEET]["A2"].value == "PROD"

    def test_an_empty_debt_still_produces_both_tabs(self, tmp_path: Path) -> None:
        target = ExcelReportWriter().write_debt_workbook(
            [], [], tmp_path / DEBT_WORKBOOK_NAME
        )

        workbook = openpyxl.load_workbook(target)
        assert workbook.sheetnames == [TECHNICAL_DEBT_SHEET, DEVIATIONS_SHEET]
        assert workbook[TECHNICAL_DEBT_SHEET].max_row == 1

    def test_a_run_writes_the_workbook_next_to_the_other_excels(
        self, tmp_path: Path
    ) -> None:
        source = tmp_path / "source"
        source.mkdir()
        generator = SalesforceDocumentationGenerator(source, tmp_path / "out")
        generator.alias = "MHINT"
        excel_dir = tmp_path / "out" / "excel"
        result = GenerationResult()

        generator._generate_excels(
            _snapshot(tmp_path), ExcelReportWriter(), excel_dir, result
        )

        assert result.debt_excel == excel_dir / DEBT_WORKBOOK_NAME
        workbook = openpyxl.load_workbook(result.debt_excel)
        assert workbook[TECHNICAL_DEBT_SHEET]["A2"].value == _DEBT_ITEM.label


class _FakeTaskManager:
    def __init__(self) -> None:
        self.result = None

    def queue_log(self, message: str) -> None:
        pass

    def start_task(self, *, status_text, task, success_message, on_success=None, notify=True):
        self.result = task()


class TestDebtScreenExport:
    def test_the_screen_rows_are_exported_without_a_full_run(
        self, tmp_path: Path
    ) -> None:
        output = tmp_path / "output"
        app = SimpleNamespace(
            task_manager=_FakeTaskManager(),
            _t=lambda key, **kwargs: key,
            _validate_output_dir=lambda: output,
        )

        target = export_debt_workbook(
            app,
            [("MHINT", {"label": "Trigger sans handler", "date_creation": "2026-01-05"})],
            [("PROD", {"label": "Champ en dur", "explanation": "Valide en comite"})],
        )

        assert target == output / "excel" / DEBT_WORKBOOK_NAME
        workbook = openpyxl.load_workbook(target)
        assert workbook[TECHNICAL_DEBT_SHEET]["A2"].value == "MHINT"
        assert workbook[TECHNICAL_DEBT_SHEET]["B2"].value == "Trigger sans handler"
        assert workbook[DEVIATIONS_SHEET]["D2"].value == "Valide en comite"
