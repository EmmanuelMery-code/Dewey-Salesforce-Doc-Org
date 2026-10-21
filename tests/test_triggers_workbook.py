"""Tests: trigger.xlsx workbook (ExcelReportWriter.write_triggers_workbook).

Contracts tested:
  - One row per object carrying at least one Apex trigger or record-triggered
    flow, sorted by API name; the label comes from the object metadata and
    stays empty when unknown.
  - Each (timing, event) column lists the trigger / flow names, one per line.
  - A "Create or Update" flow fills both the Create and Update columns.
  - Inactive triggers and non-active flows are listed with "(inactif)".
  - Flows that are not record-triggered are ignored.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.core.models import ApexArtifact, FlowInfo, ObjectInfo
from src.parsers.salesforce_parser import SalesforceMetadataParser
from src.reporting.excel_writer import ExcelReportWriter
from src.reporting.excel_writer_triggers import apex_trigger_events

HEADERS = (
    "Objet (API Name)",
    "Label",
    "Apex Before Insert",
    "Apex Before Update",
    "Apex Before Delete",
    "Apex After Insert",
    "Apex After Update",
    "Apex After Delete",
    "Apex After Undelete",
    "Flow Before Create",
    "Flow Before Update",
    "Flow Before Delete",
    "Flow After Create",
    "Flow After Update",
)


def _trigger(name: str, header: str, status: str = "Active") -> ApexArtifact:
    return ApexArtifact(
        name=name,
        kind="trigger",
        body=f"{header} {{\n    Handler.run();\n}}\n",
        source_path=Path(f"{name}.trigger"),
        status=status,
    )


def _flow(name: str, obj: str, trigger_type: str, record_trigger_type: str = "", status: str = "Active") -> FlowInfo:
    return FlowInfo(
        name=name,
        process_type="AutoLaunchedFlow",
        status=status,
        trigger_type=trigger_type,
        record_trigger_type=record_trigger_type,
        start_object=obj,
    )


def _write(tmp_path: Path, apex=(), flows=(), objects=()):
    output = ExcelReportWriter().write_triggers_workbook(
        list(apex), list(flows), list(objects), tmp_path / "trigger.xlsx"
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Triggers"]
    rows = list(workbook.active.iter_rows(values_only=True))
    return [dict(zip(HEADERS, row)) for row in rows[1:]], rows[0]


class TestApexTriggerEvents:
    def test_header_is_parsed_case_insensitively(self) -> None:
        body = "// trigger Fake on Lead (after delete)\nTrigger AccountTrg ON Account (Before Insert,  AFTER   update) {}"

        assert apex_trigger_events(body) == ("Account", ["before insert", "after update"])

    def test_body_without_header_yields_nothing(self) -> None:
        assert apex_trigger_events("public class Foo {}") == ("", [])


class TestTriggersWorkbook:
    def test_headers_and_one_row_per_object(self, tmp_path: Path) -> None:
        rows, headers = _write(
            tmp_path,
            apex=[
                _trigger("AccountTrigger", "trigger AccountTrigger on Account (before insert, after update)"),
                _trigger("CaseTrigger", "trigger CaseTrigger on Case (after undelete)"),
            ],
            flows=[_flow("Opp_BS", "Opportunity", "RecordBeforeSave", "Update")],
            objects=[ObjectInfo(api_name="Account", label="Compte")],
        )

        assert headers == HEADERS
        assert [row["Objet (API Name)"] for row in rows] == ["Account", "Case", "Opportunity"]
        account, case, opportunity = rows
        assert account["Label"] == "Compte"
        assert account["Apex Before Insert"] == "AccountTrigger"
        assert account["Apex After Update"] == "AccountTrigger"
        assert account["Apex Before Update"] is None
        assert case["Label"] is None
        assert case["Apex After Undelete"] == "CaseTrigger"
        assert opportunity["Flow Before Update"] == "Opp_BS"

    def test_several_components_are_listed_one_per_line(self, tmp_path: Path) -> None:
        rows, _ = _write(
            tmp_path,
            apex=[
                _trigger("Zeta", "trigger Zeta on Account (before insert)"),
                _trigger("Alpha", "trigger Alpha on Account (before insert)"),
            ],
        )

        assert rows[0]["Apex Before Insert"] == "Alpha\nZeta"

    def test_create_and_update_flow_fills_both_columns(self, tmp_path: Path) -> None:
        rows, _ = _write(
            tmp_path,
            flows=[
                _flow("Acc_AS", "Account", "RecordAfterSave", "CreateAndUpdate"),
                _flow("Acc_BD", "Account", "RecordBeforeDelete"),
            ],
        )

        assert rows[0]["Flow After Create"] == "Acc_AS"
        assert rows[0]["Flow After Update"] == "Acc_AS"
        assert rows[0]["Flow Before Delete"] == "Acc_BD"
        assert rows[0]["Flow Before Create"] is None

    def test_inactive_components_are_flagged(self, tmp_path: Path) -> None:
        rows, _ = _write(
            tmp_path,
            apex=[_trigger("OldTrg", "trigger OldTrg on Lead (after insert)", status="Inactive")],
            flows=[_flow("Draft_Flow", "Lead", "RecordAfterSave", "Create", status="Draft")],
        )

        assert rows[0]["Apex After Insert"] == "OldTrg (inactif)"
        assert rows[0]["Flow After Create"] == "Draft_Flow (inactif)"

    def test_non_record_triggered_flows_are_ignored(self, tmp_path: Path) -> None:
        rows, _ = _write(
            tmp_path,
            flows=[
                _flow("Screen", "", ""),
                _flow("Scheduled", "Account", "Scheduled"),
                _flow("Event", "MyEvent__e", "PlatformEvent"),
            ],
        )

        assert rows == []


FLOW_META = """<?xml version="1.0" encoding="UTF-8"?>
<Flow xmlns="http://soap.sforce.com/2006/04/metadata">
    <label>Maj Compte</label>
    <processType>AutoLaunchedFlow</processType>
    <status>Active</status>
    <start>
        <object>Account</object>
        <recordTriggerType>CreateAndUpdate</recordTriggerType>
        <triggerType>RecordBeforeSave</triggerType>
    </start>
</Flow>
"""

TRIGGER_META = """<?xml version="1.0" encoding="UTF-8"?>
<ApexTrigger xmlns="http://soap.sforce.com/2006/04/metadata">
    <apiVersion>61.0</apiVersion>
    <status>Inactive</status>
</ApexTrigger>
"""


class TestThroughTheFullParser:
    def test_parsed_source_feeds_the_workbook(self, tmp_path: Path) -> None:
        source = tmp_path / "source"
        (source / "flows").mkdir(parents=True)
        (source / "flows" / "Maj_Compte.flow-meta.xml").write_text(FLOW_META, encoding="utf-8")
        (source / "triggers").mkdir()
        (source / "triggers" / "AccountTrigger.trigger").write_text(
            "trigger AccountTrigger on Account (after delete) {}\n", encoding="utf-8"
        )
        (source / "triggers" / "AccountTrigger.trigger-meta.xml").write_text(
            TRIGGER_META, encoding="utf-8"
        )

        snapshot = SalesforceMetadataParser(source).parse()
        assert snapshot.flows[0].record_trigger_type == "CreateAndUpdate"

        rows, _ = _write(tmp_path, apex=snapshot.apex_artifacts, flows=snapshot.flows)

        assert rows[0]["Flow Before Create"] == "Maj_Compte"
        assert rows[0]["Flow Before Update"] == "Maj_Compte"
        assert rows[0]["Apex After Delete"] == "AccountTrigger (inactif)"
