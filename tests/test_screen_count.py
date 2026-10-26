"""Comptage des écrans : squads des permission sets et du data dictionary."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.core.models.security import ObjectPermission, SecurityArtifact
from src.reporting.excel_writer_screens import write_screen_count_workbook
from src.reporting.screen_count_collect import (
    collect_screens,
    grants_from_permission_sets,
    split_squads,
    squad_from_description,
)
from src.ui.translations import TRANSLATIONS

NS = 'xmlns="http://soap.sforce.com/2006/04/metadata"'


def test_squad_phrases_and_platform_objects_are_normalised() -> None:
    assert squad_from_description(
        "droits sur les objets de la squad parcours de Vente aux Super Users"
    ) == "Parcours de vente"
    assert squad_from_description("perimetre Parcours de vente.") == "Parcours de vente"
    assert squad_from_description("périmètre Vision Client") == "Vision Client"
    assert split_squads("Tous | Vision Client") == ["Vision Client"]
    assert split_squads("Tous") == []
    assert split_squads("Vision Client | Parcours de vente") == [
        "Vision Client",
        "Parcours de vente",
    ]

    grants = grants_from_permission_sets(
        [
            SecurityArtifact(
                name="PS_PV",
                kind="permissionset",
                description="squad parcours de Vente aux Super Users",
                object_permissions=[
                    ObjectPermission("Opportunity", allow_read=True),
                    ObjectPermission("Account", allow_read=True),
                    ObjectPermission("SF_Async__c", allow_read=True),
                ],
            ),
            SecurityArtifact(
                name="PS_VC",
                kind="permissionset",
                description="périmètre Vision Client",
                object_permissions=[
                    ObjectPermission("Account", allow_read=True),
                    ObjectPermission("Contact", allow_read=True),
                ],
            ),
        ]
    )
    assert grants["Parcours de vente"] == {"Opportunity", "Account"}
    assert "SF_Async__c" not in grants["Parcours de vente"]
    assert grants["Vision Client"] == {"Account", "Contact"}


def test_workbook_counts_screens_with_both_squad_methods(tmp_path: Path) -> None:
    source = _build_package(tmp_path)
    result = collect_screens(
        [source],
        squad_grants={
            "Parcours de vente": {"Opportunity", "Account"},
            "Vision Client": {"Account", "Contact"},
        },
        known_objects={"Opportunity", "Account", "Contact", "InsurancePolicy"},
        object_squad={
            "Account": "Vision Client",
            "Contact": "Vision Client",
            "Opportunity": "Parcours de vente",
            "InsurancePolicy": "Vision Client | Parcours de vente",
        },
        object_piloted_by={"Account": "RCU", "Opportunity": "Salesforce"},
        object_squad_consumer={"Account": "Parcours de vente"},
    )
    path = write_screen_count_workbook(result, tmp_path / "ecrans.xlsx")
    workbook = load_workbook(path)

    assert workbook.sheetnames == [
        "Synthese",
        "Synthese dictionnaire",
        "Detail dictionnaire",
        "Detail",
        "Hors perimetre",
        "Objets",
        "Methode",
    ]
    summary = workbook["Synthese"]
    assert [summary.cell(4, col).value for col in range(1, 7)] == [
        "Type d'écran",
        "Parcours de vente",
        "Vision Client",
        "Objets partagés",
        "PV et VC",
        "Total",
    ]
    assert _row(summary, "Lightning Page") == [2, 0, 1, 0, 3]
    assert _row(summary, "Screen Flow") == [2, 0, 1, 1, 4]
    assert _row(summary, "OmniScript") == [3, 0, 0, 0, 3]
    assert _row(summary, "LWC") == [1, 1, 0, 0, 2]
    assert _row(summary, "Total") == [8, 1, 2, 1, 12]
    assert summary._charts

    dictionary = workbook["Synthese dictionnaire"]
    assert [dictionary.cell(4, col).value for col in range(1, 7)] == [
        "Type d'écran",
        "Parcours de vente",
        "Vision Client",
        "Plusieurs squads",
        "Non renseigné",
        "Total",
    ]
    assert _row(dictionary, "Lightning Page") == [1, 1, 1, 1, 4]
    assert _row(dictionary, "Screen Flow") == [2, 1, 1, 0, 4]
    assert _row(dictionary, "OmniScript") == [3, 0, 0, 0, 3]
    assert _row(dictionary, "LWC") == [1, 1, 0, 1, 3]
    assert _row(dictionary, "Total") == [7, 3, 2, 2, 14]

    details = _indexed(workbook["Detail dictionnaire"])
    assert details["sfSocle"][4] == "Parcours de vente"
    assert details["Account_Record"][4] == "Vision Client"
    assert details["Policy_Record"][4] == "Plusieurs squads"
    assert details["OLI_Record"][4] == "Non renseigné"
    assert not details["noiseLwc"][3]

    outside = {row[1] for row in workbook["Hors perimetre"].iter_rows(min_row=2, values_only=True)}
    assert {"Util_Bar", "sfDraft", "internalLwc", "utilityLwc", "headlessLwc", "PV_Demo_1"} <= outside
    assert "sfAuto" not in outside


def test_menu_entry_follows_the_data_dictionary() -> None:
    assert TRANSLATIONS["fr"]["menu_count_screens"] == "Compter les écrans"
    assert TRANSLATIONS["en"]["menu_count_screens"] == "Count screens"
    menu = (Path(__file__).resolve().parents[1] / "src/ui/app_ui_menu_mixin.py").read_text(encoding="utf-8")
    dictionary = menu.index("menu_create_data_dictionary")
    screens = menu.index("menu_count_screens")
    findings = menu.index("menu_create_findings_document")
    assert dictionary < screens < findings


def _row(worksheet, label: str) -> list[int]:
    for row in worksheet.iter_rows(min_row=5, max_row=9, values_only=True):
        if row[0] == label:
            return list(row[1:])
    raise AssertionError(label)


def _indexed(worksheet) -> dict[str, tuple]:
    return {row[1]: row for row in worksheet.iter_rows(min_row=2, values_only=True)}


def _build_package(tmp_path: Path) -> Path:
    root = tmp_path / "force-app"
    _write(root / "flexipages" / "Opp_Record.flexipage-meta.xml", _page("Opportunity", "RecordPage", "c:oppLwc"))
    _write(root / "flexipages" / "Account_Record.flexipage-meta.xml", _page("Account", "RecordPage"))
    _write(root / "flexipages" / "OLI_Record.flexipage-meta.xml", _page("OpportunityLineItem", "RecordPage"))
    _write(root / "flexipages" / "Policy_Record.flexipage-meta.xml", _page("InsurancePolicy", "RecordPage"))
    _write(root / "flexipages" / "Util_Bar.flexipage-meta.xml", _page("", "UtilityBar"))
    _write(root / "flexipages" / "Home_Page.flexipage-meta.xml", _page("", "HomePage"))

    _write(root / "flows" / "sfOpp.flow-meta.xml", _flow("Active", "Flow", "Opportunity"))
    _write(root / "flows" / "sfAccount.flow-meta.xml", _flow("Active", "Flow", "Account"))
    _write(root / "flows" / "sfBoth.flow-meta.xml", _flow("Active", "Flow", "Opportunity", "Contact"))
    _write(root / "flows" / "sfSocle.flow-meta.xml", _flow("Active", "Flow", "Account", "Opportunity"))
    _write(root / "flows" / "sfAuto.flow-meta.xml", _flow("Active", "AutoLaunchedFlow", "Opportunity"))
    _write(root / "flows" / "sfDraft.flow-meta.xml", _flow("Draft", "Flow", "Opportunity"))

    _write(root / "omniDataTransforms" / "DR_Opp_1.rpt-meta.xml", _dr("Opportunity", "Turbo Extract"))
    _write(root / "omniDataTransforms" / "DR_Old_1.rpt-meta.xml", _dr("Lead", "json"))
    _write(
        root / "omniIntegrationProcedures" / "PV_Create_English_1.oip-meta.xml",
        _procedure("DR_Opp"),
    )
    _write(root / "omniScripts" / "PV_Demo_fr_1.os-meta.xml", _omni(1, "DR_Old", ""))
    _write(root / "omniScripts" / "PV_Demo_fr_2.os-meta.xml", _omni(2, "DR_Opp", ""))
    _write(root / "omniScripts" / "EV_Host_fr_1.os-meta.xml", _omni(1, "", "", embed="PV|Demo|fr", os_type="EV", subtype="Host"))
    _write(root / "omniScripts" / "PV_ViaIp_fr_1.os-meta.xml", _omni(1, "", "PV_Create", os_type="PV", subtype="ViaIp"))
    _write(root / "omniScripts" / "PV_NoStep_fr_1.os-meta.xml", _omni(1, "DR_Opp", "", subtype="NoStep", step=False))

    _lwc(root, "oppLwc", exposed=True, targets=["lightning__RecordPage"], schema="Opportunity")
    _lwc(root, "contactLwc", exposed=True, targets=["lightning__RecordAction"], action="ScreenAction", sobject="Contact")
    _lwc(root, "noiseLwc", exposed=True, targets=["lightning__RecordPage"], comment="{Event}")
    _lwc(root, "internalLwc", exposed=False, targets=["lightning__RecordPage"])
    _lwc(root, "utilityLwc", exposed=True, targets=["lightning__UtilityBar"])
    _lwc(root, "headlessLwc", exposed=True, targets=["lightning__RecordAction"], action="Action")
    return root


def _write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def _page(sobject: str, kind: str, component: str = "") -> str:
    component_xml = (
        f"<flexiPageRegions><itemInstances><componentInstance><componentName>{component}</componentName></componentInstance></itemInstances></flexiPageRegions>"
        if component
        else ""
    )
    sobject_xml = f"<sobjectType>{sobject}</sobjectType>" if sobject else ""
    return f"<FlexiPage {NS}><masterLabel>{kind}</masterLabel><type>{kind}</type>{sobject_xml}{component_xml}</FlexiPage>"


def _flow(status: str, process: str, *objects: str) -> str:
    lookups = "".join(f"<recordLookups><object>{name}</object></recordLookups>" for name in objects)
    return (
        f"<Flow {NS}><label>Flow</label><processType>{process}</processType>"
        f"<status>{status}</status><screens><name>Ecran</name></screens>{lookups}</Flow>"
    )


def _dr(source: str, target: str) -> str:
    return (
        f"<OmniDataTransform {NS}><inputObjectName>{source}</inputObjectName>"
        f"<outputObjectName>{target}</outputObjectName></OmniDataTransform>"
    )


def _procedure(bundle: str) -> str:
    config = f"&quot;bundle&quot; : &quot;{bundle}&quot;"
    return (
        f"<OmniIntegrationProcedure {NS}><isActive>true</isActive><isIntegrationProcedure>true</isIntegrationProcedure>"
        f"<propertySetConfig>{config}</propertySetConfig><subType>Create</subType><type>PV</type>"
        f"<versionNumber>1.0</versionNumber></OmniIntegrationProcedure>"
    )


def _omni(
    version: int,
    bundle: str,
    procedure: str,
    *,
    embed: str = "",
    os_type: str = "PV",
    subtype: str = "Demo",
    step: bool = True,
) -> str:
    parts = []
    if bundle:
        parts.append(f"&quot;bundle&quot; : &quot;{bundle}&quot;")
    if procedure:
        parts.append(f"&quot;integrationProcedureKey&quot; : &quot;{procedure}&quot;")
    config = ", ".join(parts)
    step_xml = "<omniProcessElements><type>Step</type></omniProcessElements>" if step else ""
    embed_xml = f"<embeddedOmniScriptKey>{embed}</embeddedOmniScriptKey>" if embed else ""
    return (
        f"<OmniScript {NS}><isActive>true</isActive><isIntegrationProcedure>false</isIntegrationProcedure>"
        f"<language>fr</language><name>{subtype}</name>{step_xml}{embed_xml}"
        f"<propertySetConfig>{config}</propertySetConfig><subType>{subtype}</subType><type>{os_type}</type>"
        f"<versionNumber>{version}.0</versionNumber></OmniScript>"
    )


def _lwc(
    root: Path,
    name: str,
    *,
    exposed: bool,
    targets: list[str],
    schema: str = "",
    action: str = "",
    sobject: str = "",
    comment: str = "",
) -> None:
    target_xml = "".join(f"<target>{target}</target>" for target in targets)
    config = ""
    if action or sobject:
        declared = " ".join(targets)
        object_xml = f"<objects><object>{sobject}</object></objects>" if sobject else ""
        action_xml = f"<actionType>{action}</actionType>" if action else ""
        config = f'<targetConfigs><targetConfig targets="{declared}">{action_xml}{object_xml}</targetConfig></targetConfigs>'
    _write(
        root / "lwc" / name / f"{name}.js-meta.xml",
        f"<LightningComponentBundle {NS}><isExposed>{str(exposed).lower()}</isExposed>"
        f"<masterLabel>{name}</masterLabel><targets>{target_xml}</targets>{config}</LightningComponentBundle>",
    )
    body = ""
    if schema:
        body += f"import FIELD from '@salesforce/schema/{schema}.Name';\n"
    if comment:
        body += f"/** @param {comment} event */\n"
    _write(root / "lwc" / name / f"{name}.js", body)
